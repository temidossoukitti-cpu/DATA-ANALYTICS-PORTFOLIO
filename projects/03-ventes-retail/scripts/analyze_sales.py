"""Build the business analysis, KPI exports, charts, and final report."""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/matplotlib-retail-sales")

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = PROJECT_ROOT / "data" / "processed" / "superstore_clean.csv"
KPI_PATH = PROJECT_ROOT / "data" / "processed" / "business_kpis.csv"
EXCEL_PATH = PROJECT_ROOT / "reports" / "retail_sales_analysis.xlsx"
REPORT_PATH = PROJECT_ROOT / "reports" / "business_analysis.md"
VISUALS_DIR = PROJECT_ROOT / "visuals"

DATE_COLUMNS = ["Order Date", "Ship Date"]
COLORS = {
    "navy": "#17324D",
    "blue": "#2878B5",
    "teal": "#2A9D8F",
    "green": "#4C956C",
    "red": "#C44536",
    "orange": "#E09F3E",
    "gray": "#6B7280",
    "light": "#D9E2EC",
}


def load_data() -> pd.DataFrame:
    """Load the cleaned dataset with its persisted schema."""
    if not DATA_PATH.exists():
        raise FileNotFoundError(
            f"Clean dataset not found: {DATA_PATH}. Run scripts/clean_data.py first."
        )

    data = pd.read_csv(
        DATA_PATH,
        parse_dates=DATE_COLUMNS,
        dtype={"Postal Code": "string"},
    )
    required = {
        "Order ID",
        "Order Date",
        "Ship Date",
        "Customer ID",
        "Region",
        "Category",
        "Sub-Category",
        "Sales",
        "Quantity",
        "Discount",
        "Profit",
        "Potential Duplicate",
    }
    missing = required.difference(data.columns)
    if missing:
        raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")

    data["Year"] = data["Order Date"].dt.year
    data["Month"] = data["Order Date"].dt.to_period("M").dt.to_timestamp()
    data["Shipping Days"] = (data["Ship Date"] - data["Order Date"]).dt.days
    data["Profit Margin"] = np.where(
        data["Sales"].ne(0),
        data["Profit"] / data["Sales"],
        np.nan,
    )
    data["Discount Band"] = pd.cut(
        data["Discount"],
        bins=[-0.001, 0, 0.2, 0.4, 0.6, 1.0],
        labels=["0%", "1-20%", "21-40%", "41-60%", ">60%"],
    )
    return data


def aggregate_performance(data: pd.DataFrame, group_columns: list[str]) -> pd.DataFrame:
    """Aggregate sales, profit, volume, and margin for business dimensions."""
    result = (
        data.groupby(group_columns, observed=True)
        .agg(
            Sales=("Sales", "sum"),
            Profit=("Profit", "sum"),
            Quantity=("Quantity", "sum"),
            Orders=("Order ID", "nunique"),
            Average_Discount=("Discount", "mean"),
        )
        .reset_index()
    )
    result["Profit_Margin"] = np.where(
        result["Sales"].ne(0),
        result["Profit"] / result["Sales"],
        np.nan,
    )
    return result


def build_analysis_tables(data: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Create the detailed tables used by exports, charts, and reporting."""
    yearly = aggregate_performance(data, ["Year"]).sort_values("Year")
    yearly["Sales_Growth"] = yearly["Sales"].pct_change()
    yearly["Profit_Growth"] = yearly["Profit"].pct_change()

    monthly = aggregate_performance(data, ["Month"]).sort_values("Month")
    category = aggregate_performance(data, ["Category"]).sort_values(
        "Sales", ascending=False
    )
    subcategory = aggregate_performance(data, ["Sub-Category"]).sort_values(
        "Profit", ascending=False
    )
    region = aggregate_performance(data, ["Region"]).sort_values(
        "Sales", ascending=False
    )
    segment = aggregate_performance(data, ["Segment"]).sort_values(
        "Sales", ascending=False
    )
    state = aggregate_performance(data, ["State"]).sort_values(
        "Profit", ascending=False
    )
    discount = aggregate_performance(data, ["Discount Band"])
    discount["Discount Band"] = discount["Discount Band"].astype("string")

    customers = aggregate_performance(
        data,
        ["Customer ID", "Customer Name"],
    ).sort_values("Sales", ascending=False)

    products = aggregate_performance(
        data,
        ["Product ID", "Product Name", "Category", "Sub-Category"],
    ).sort_values("Profit", ascending=False)

    shipping = (
        data.groupby("Ship Mode", observed=True)
        .agg(
            Orders=("Order ID", "nunique"),
            Sales=("Sales", "sum"),
            Profit=("Profit", "sum"),
            Average_Shipping_Days=("Shipping Days", "mean"),
            Median_Shipping_Days=("Shipping Days", "median"),
        )
        .reset_index()
        .sort_values("Orders", ascending=False)
    )
    shipping["Profit_Margin"] = shipping["Profit"] / shipping["Sales"]

    return {
        "Yearly": yearly,
        "Monthly": monthly,
        "Categories": category,
        "Subcategories": subcategory,
        "Regions": region,
        "Segments": segment,
        "States": state,
        "Discounts": discount,
        "Customers": customers,
        "Products": products,
        "Shipping": shipping,
    }


def build_kpis(data: pd.DataFrame, tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Calculate executive KPIs at company level."""
    order_sales = data.groupby("Order ID")["Sales"].sum()
    customer_orders = data.groupby("Customer ID")["Order ID"].nunique()
    yearly = tables["Yearly"]
    first_sales = yearly.iloc[0]["Sales"]
    last_sales = yearly.iloc[-1]["Sales"]
    periods = max(len(yearly) - 1, 1)

    values = [
        ("Total Sales", data["Sales"].sum(), "currency"),
        ("Total Profit", data["Profit"].sum(), "currency"),
        ("Profit Margin", data["Profit"].sum() / data["Sales"].sum(), "percent"),
        ("Orders", data["Order ID"].nunique(), "integer"),
        ("Customers", data["Customer ID"].nunique(), "integer"),
        ("Products", data["Product ID"].nunique(), "integer"),
        ("Units Sold", data["Quantity"].sum(), "integer"),
        ("Average Order Value", order_sales.mean(), "currency"),
        ("Average Discount", data["Discount"].mean(), "percent"),
        ("Average Shipping Days", data["Shipping Days"].mean(), "decimal"),
        (
            "Repeat Customer Rate",
            customer_orders.gt(1).mean(),
            "percent",
        ),
        (
            "Loss-making Sales Share",
            data.loc[data["Profit"].lt(0), "Sales"].sum() / data["Sales"].sum(),
            "percent",
        ),
        (
            "Sales CAGR",
            (last_sales / first_sales) ** (1 / periods) - 1,
            "percent",
        ),
        (
            "Latest Annual Sales Growth",
            yearly.iloc[-1]["Sales_Growth"],
            "percent",
        ),
        (
            "Latest Annual Profit Growth",
            yearly.iloc[-1]["Profit_Growth"],
            "percent",
        ),
    ]
    return pd.DataFrame(values, columns=["KPI", "Value", "Format"])


def currency_axis(value: float, _: int) -> str:
    """Format chart values as compact US dollars."""
    if abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"${value / 1_000:.0f}K"
    return f"${value:.0f}"


def style_axis(axis: plt.Axes, title: str, ylabel: str = "") -> None:
    """Apply a consistent, restrained chart style."""
    axis.set_title(title, loc="left", fontsize=14, fontweight="bold", pad=14)
    axis.set_ylabel(ylabel)
    axis.spines[["top", "right"]].set_visible(False)
    axis.grid(axis="y", color="#E5E7EB", linewidth=0.8)
    axis.set_axisbelow(True)


def save_figure(figure: plt.Figure, filename: str) -> None:
    """Persist a chart and release its Matplotlib resources."""
    VISUALS_DIR.mkdir(parents=True, exist_ok=True)
    figure.tight_layout()
    figure.savefig(VISUALS_DIR / filename, dpi=160, bbox_inches="tight")
    plt.close(figure)


def create_charts(tables: dict[str, pd.DataFrame]) -> None:
    """Generate charts, with each visual tied to a business question."""
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.labelcolor": COLORS["navy"],
            "xtick.color": COLORS["gray"],
            "ytick.color": COLORS["gray"],
            "text.color": COLORS["navy"],
        }
    )

    yearly = tables["Yearly"]
    figure, axes = plt.subplots(1, 2, figsize=(12, 4.5))
    axes[0].bar(yearly["Year"].astype(str), yearly["Sales"], color=COLORS["blue"])
    axes[1].bar(yearly["Year"].astype(str), yearly["Profit"], color=COLORS["green"])
    style_axis(axes[0], "How has annual revenue evolved?", "Sales")
    style_axis(axes[1], "How has annual profit evolved?", "Profit")
    axes[0].yaxis.set_major_formatter(FuncFormatter(currency_axis))
    axes[1].yaxis.set_major_formatter(FuncFormatter(currency_axis))
    save_figure(figure, "01_annual_performance.png")

    monthly = tables["Monthly"]
    figure, axis = plt.subplots(figsize=(12, 4.8))
    axis.plot(
        monthly["Month"],
        monthly["Sales"],
        color=COLORS["blue"],
        linewidth=2,
        label="Monthly sales",
    )
    rolling_sales = monthly["Sales"].rolling(3, min_periods=1).mean()
    axis.plot(
        monthly["Month"],
        rolling_sales,
        color=COLORS["orange"],
        linewidth=2,
        label="3-month average",
    )
    style_axis(axis, "What is the monthly sales trend?", "Sales")
    axis.yaxis.set_major_formatter(FuncFormatter(currency_axis))
    axis.legend(frameon=False)
    save_figure(figure, "02_monthly_sales_trend.png")

    category = tables["Categories"].sort_values("Sales")
    figure, axes = plt.subplots(1, 2, figsize=(12, 4.5))
    axes[0].barh(category["Category"], category["Sales"], color=COLORS["blue"])
    axes[1].barh(category["Category"], category["Profit"], color=COLORS["green"])
    style_axis(axes[0], "Which categories generate revenue?", "Category")
    style_axis(axes[1], "Which categories generate profit?", "Category")
    axes[0].xaxis.set_major_formatter(FuncFormatter(currency_axis))
    axes[1].xaxis.set_major_formatter(FuncFormatter(currency_axis))
    save_figure(figure, "03_category_performance.png")

    subcategory = tables["Subcategories"].sort_values("Profit")
    bar_colors = np.where(
        subcategory["Profit"].ge(0),
        COLORS["green"],
        COLORS["red"],
    )
    figure, axis = plt.subplots(figsize=(10, 7))
    axis.barh(subcategory["Sub-Category"], subcategory["Profit"], color=bar_colors)
    axis.axvline(0, color=COLORS["navy"], linewidth=0.8)
    style_axis(
        axis,
        "Which sub-categories create or destroy profit?",
        "Sub-category",
    )
    axis.xaxis.set_major_formatter(FuncFormatter(currency_axis))
    save_figure(figure, "04_subcategory_profit.png")

    region = tables["Regions"].sort_values("Sales")
    figure, axes = plt.subplots(1, 2, figsize=(12, 4.8))
    axes[0].barh(region["Region"], region["Sales"], color=COLORS["blue"])
    margin_colors = np.where(
        region["Profit_Margin"].ge(0.1),
        COLORS["green"],
        COLORS["orange"],
    )
    axes[1].barh(region["Region"], region["Profit_Margin"], color=margin_colors)
    style_axis(axes[0], "Where is revenue generated?", "Region")
    style_axis(axes[1], "Which regions are most profitable?", "Region")
    axes[0].xaxis.set_major_formatter(FuncFormatter(currency_axis))
    axes[1].xaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x:.0%}"))
    save_figure(figure, "05_region_performance.png")

    discount = tables["Discounts"]
    figure, axis = plt.subplots(figsize=(9, 4.8))
    discount_colors = np.where(
        discount["Profit_Margin"].ge(0),
        COLORS["teal"],
        COLORS["red"],
    )
    axis.bar(
        discount["Discount Band"],
        discount["Profit_Margin"],
        color=discount_colors,
    )
    axis.axhline(0, color=COLORS["navy"], linewidth=0.8)
    style_axis(
        axis,
        "How do discount levels affect profit margin?",
        "Profit margin",
    )
    axis.yaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x:.0%}"))
    save_figure(figure, "06_discount_profitability.png")

    states = tables["States"]
    selected_states = pd.concat([states.head(10), states.tail(10)]).sort_values("Profit")
    state_colors = np.where(
        selected_states["Profit"].ge(0),
        COLORS["green"],
        COLORS["red"],
    )
    figure, axis = plt.subplots(figsize=(10, 7.5))
    axis.barh(selected_states["State"], selected_states["Profit"], color=state_colors)
    axis.axvline(0, color=COLORS["navy"], linewidth=0.8)
    style_axis(
        axis,
        "Which states contribute most and least to profit?",
        "State",
    )
    axis.xaxis.set_major_formatter(FuncFormatter(currency_axis))
    save_figure(figure, "07_state_profit_extremes.png")

    customers = tables["Customers"].head(10).sort_values("Sales")
    figure, axis = plt.subplots(figsize=(10, 5.5))
    axis.barh(customers["Customer Name"], customers["Sales"], color=COLORS["blue"])
    style_axis(axis, "Who are the top 10 customers by revenue?", "Customer")
    axis.xaxis.set_major_formatter(FuncFormatter(currency_axis))
    save_figure(figure, "08_top_customers.png")


def export_excel(
    kpis: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
) -> None:
    """Export KPI and analysis tables to a formatted Excel workbook."""
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl") as writer:
        kpis.to_excel(writer, sheet_name="Executive KPIs", index=False)
        for sheet_name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_name[:31], index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="17324D")
        header_font = Font(color="FFFFFF", bold=True)

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for column_cells in worksheet.columns:
                values = [str(cell.value or "") for cell in column_cells]
                width = min(max(len(value) for value in values) + 2, 45)
                worksheet.column_dimensions[
                    get_column_letter(column_cells[0].column)
                ].width = width

                header = str(column_cells[0].value or "")
                if header in {"Sales", "Profit", "Average Order Value"}:
                    for cell in column_cells[1:]:
                        cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
                elif "Margin" in header or "Growth" in header or "Discount" in header:
                    for cell in column_cells[1:]:
                        cell.number_format = "0.0%"
                elif header in {"Order Date", "Ship Date", "Month"}:
                    for cell in column_cells[1:]:
                        cell.number_format = "yyyy-mm-dd"

            if worksheet.title == "Executive KPIs":
                for row in range(2, worksheet.max_row + 1):
                    value_cell = worksheet.cell(row=row, column=2)
                    value_format = worksheet.cell(row=row, column=3).value
                    if value_format == "currency":
                        value_cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
                    elif value_format == "percent":
                        value_cell.number_format = "0.0%"
                    elif value_format == "integer":
                        value_cell.number_format = "#,##0"
                    else:
                        value_cell.number_format = "0.00"


def format_kpi(value: float, value_format: str) -> str:
    """Format one KPI for the Markdown report."""
    if value_format == "currency":
        return f"${value:,.2f}"
    if value_format == "percent":
        return f"{value:.1%}"
    if value_format == "integer":
        return f"{value:,.0f}"
    return f"{value:,.2f}"


def markdown_table(
    data: pd.DataFrame,
    columns: list[str],
    formats: dict[str, str],
) -> str:
    """Render selected DataFrame columns as a compact Markdown table."""
    headers = [column.replace("_", " ") for column in columns]
    lines = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(":--" if i == 0 else "--:" for i in range(len(columns))) + "|",
    ]

    for _, row in data.iterrows():
        values = []
        for column in columns:
            value = row[column]
            value_format = formats.get(column)
            if pd.isna(value):
                values.append("N/A")
            elif value_format == "currency":
                values.append(f"${value:,.0f}")
            elif value_format == "percent":
                values.append(f"{value:.1%}")
            elif value_format == "decimal":
                values.append(f"{value:,.2f}")
            elif value_format == "integer":
                values.append(f"{value:,.0f}")
            elif value_format == "year":
                values.append(f"{value:.0f}")
            else:
                values.append(str(value))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def build_report(
    data: pd.DataFrame,
    kpis: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
) -> str:
    """Build the executive business report from calculated results."""
    kpi_values = kpis.set_index("KPI")["Value"]
    yearly = tables["Yearly"]
    categories = tables["Categories"]
    subcategories = tables["Subcategories"]
    regions = tables["Regions"]
    states = tables["States"]
    discounts = tables["Discounts"]
    customers = tables["Customers"]
    products = tables["Products"]

    best_year = yearly.loc[yearly["Sales"].idxmax()]
    best_category = categories.loc[categories["Profit"].idxmax()]
    weakest_category = categories.loc[categories["Profit_Margin"].idxmin()]
    best_region = regions.loc[regions["Profit"].idxmax()]
    worst_subcategories = subcategories.nsmallest(3, "Profit")
    best_subcategories = subcategories.nlargest(3, "Profit")
    worst_states = states.nsmallest(5, "Profit")
    high_discount = discounts.loc[discounts["Discount Band"].isin(["41-60%", ">60%"])]
    high_discount_sales = high_discount["Sales"].sum()
    high_discount_profit = high_discount["Profit"].sum()
    top_10_share = customers.head(10)["Sales"].sum() / data["Sales"].sum()
    top_product = products.iloc[0]
    worst_product = products.iloc[-1]

    kpi_rows = "\n".join(
        f"| {row.KPI} | {format_kpi(row.Value, row.Format)} |"
        for row in kpis.itertuples(index=False)
    )

    yearly_table = markdown_table(
        yearly,
        ["Year", "Sales", "Profit", "Profit_Margin", "Sales_Growth"],
        {
            "Year": "year",
            "Sales": "currency",
            "Profit": "currency",
            "Profit_Margin": "percent",
            "Sales_Growth": "percent",
        },
    )
    category_table = markdown_table(
        categories,
        ["Category", "Sales", "Profit", "Profit_Margin"],
        {
            "Sales": "currency",
            "Profit": "currency",
            "Profit_Margin": "percent",
        },
    )
    region_table = markdown_table(
        regions,
        ["Region", "Sales", "Profit", "Profit_Margin"],
        {
            "Sales": "currency",
            "Profit": "currency",
            "Profit_Margin": "percent",
        },
    )
    loss_subcategory_table = markdown_table(
        worst_subcategories,
        ["Sub-Category", "Sales", "Profit", "Profit_Margin", "Average_Discount"],
        {
            "Sales": "currency",
            "Profit": "currency",
            "Profit_Margin": "percent",
            "Average_Discount": "percent",
        },
    )
    loss_state_table = markdown_table(
        worst_states,
        ["State", "Sales", "Profit", "Profit_Margin"],
        {
            "Sales": "currency",
            "Profit": "currency",
            "Profit_Margin": "percent",
        },
    )

    return f"""# Analyse metier des ventes Retail

## 1. Synthese executive

Le jeu de donnees couvre la periode du **{data["Order Date"].min():%Y-%m-%d}**
au **{data["Order Date"].max():%Y-%m-%d}**. L'entreprise a realise
**${kpi_values["Total Sales"]:,.0f}** de chiffre d'affaires et
**${kpi_values["Total Profit"]:,.0f}** de profit, soit une marge globale de
**{kpi_values["Profit Margin"]:.1%}**.

Le chiffre d'affaires progresse sur la periode avec un taux de croissance annuel
compose de **{kpi_values["Sales CAGR"]:.1%}**. L'annee la plus forte est
**{int(best_year["Year"])}**, avec **${best_year["Sales"]:,.0f}** de ventes.
La croissance ne suffit toutefois pas a garantir la rentabilite : certaines
sous-categories, certains Etats et les remises elevees detruisent une part
significative de la marge.

## 2. Indicateurs cles

| KPI | Valeur |
|:--|--:|
{kpi_rows}

## 3. Evolution du chiffre d'affaires et du profit

{yearly_table}

![Performance annuelle](../visuals/01_annual_performance.png)

![Tendance mensuelle](../visuals/02_monthly_sales_trend.png)

Le chiffre d'affaires de la derniere annee progresse de
**{yearly.iloc[-1]["Sales_Growth"]:.1%}** par rapport a l'annee precedente.
Le profit evolue de **{yearly.iloc[-1]["Profit_Growth"]:.1%}** sur la meme
periode. Les variations mensuelles restent marquees, ce qui justifie un suivi
du budget et des stocks a une granularite mensuelle.

## 4. Performance des categories et produits

{category_table}

![Performance par categorie](../visuals/03_category_performance.png)

La categorie la plus contributrice au profit est **{best_category["Category"]}**
avec **${best_category["Profit"]:,.0f}**. La marge la plus faible est celle de
**{weakest_category["Category"]}** a **{weakest_category["Profit_Margin"]:.1%}**.

Les trois sous-categories les plus rentables sont
**{", ".join(best_subcategories["Sub-Category"].tolist())}**. A l'inverse, les
sous-categories suivantes doivent etre traitees en priorite :

{loss_subcategory_table}

![Profit par sous-categorie](../visuals/04_subcategory_profit.png)

Le produit le plus rentable est **{top_product["Product Name"]}**
(**${top_product["Profit"]:,.0f}** de profit). Le produit le moins rentable est
**{worst_product["Product Name"]}**
(**${worst_product["Profit"]:,.0f}**).

## 5. Performance geographique

{region_table}

![Performance regionale](../visuals/05_region_performance.png)

La region **{best_region["Region"]}** apporte le profit le plus eleve avec
**${best_region["Profit"]:,.0f}**. Les cinq Etats les moins rentables sont :

{loss_state_table}

![Etats les plus et moins rentables](../visuals/07_state_profit_extremes.png)

## 6. Clients

Les **10 premiers clients** representent **{top_10_share:.1%}** du chiffre
d'affaires. Le taux de clients ayant passe plus d'une commande est de
**{kpi_values["Repeat Customer Rate"]:.1%}**.

![Top clients](../visuals/08_top_customers.png)

La concentration limitee du chiffre d'affaires sur les dix premiers clients
reduit la dependance a quelques comptes, mais laisse une opportunite de mieux
segmenter et developper les clients recurrents.

## 7. Remises et rentabilite

![Impact des remises](../visuals/06_discount_profitability.png)

Les ventes accordees avec une remise superieure a 40 % representent
**${high_discount_sales:,.0f}** de chiffre d'affaires et generent
**${high_discount_profit:,.0f}** de profit. Le lien observe entre niveau de
remise et marge est descriptif ; il ne prouve pas a lui seul une causalite.

## 8. Recommandations

1. **Encadrer les remises elevees.** Mettre en place un seuil d'approbation et
   suivre la marge attendue avant toute remise superieure a 40 %.
2. **Redresser les sous-categories deficitaires.** Revoir en priorite les prix,
   couts, remises et conditions logistiques de
   **{", ".join(worst_subcategories["Sub-Category"].tolist())}**.
3. **Agir sur les zones geographiques en perte.** Construire un plan par Etat
   pour les cinq marches les moins rentables, avec objectifs de marge.
4. **Proteger les moteurs de profit.** Securiser la disponibilite et la
   promotion des sous-categories
   **{", ".join(best_subcategories["Sub-Category"].tolist())}**.
5. **Piloter la croissance par la marge.** Completer les objectifs de chiffre
   d'affaires par un objectif de profit et de marge au niveau categorie,
   sous-categorie, region et client.
6. **Developper la fidelisation.** Identifier les clients a forte valeur,
   suivre leur frequence d'achat et proposer des actions ciblees sans remise
   automatique.

## 9. Limites de l'analyse

- Les donnees couvrent uniquement les Etats-Unis et la periode disponible.
- Aucun cout detaille autre que le profit calcule n'est fourni.
- Les deux lignes marquees `Potential Duplicate` sont conservees, conformement
  a la decision de nettoyage ; leur validation metier reste necessaire.
- Les codes postaux ayant perdu un zero initial ne sont pas reconstruits.
- Les conclusions sur les remises sont des associations descriptives.

## 10. Livrables

- Donnees nettoyees : `data/processed/superstore_clean.csv`
- KPI : `data/processed/business_kpis.csv`
- Classeur d'analyse : `reports/retail_sales_analysis.xlsx`
- Graphiques : `visuals/`
"""


def save_outputs(
    data: pd.DataFrame,
    kpis: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
) -> None:
    """Write all analysis deliverables."""
    KPI_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    kpis.to_csv(KPI_PATH, index=False)
    export_excel(kpis, tables)
    create_charts(tables)
    REPORT_PATH.write_text(build_report(data, kpis, tables), encoding="utf-8")


def verify_outputs(kpis: pd.DataFrame, tables: dict[str, pd.DataFrame]) -> None:
    """Check that expected files and core accounting relationships are valid."""
    expected_files = [
        KPI_PATH,
        EXCEL_PATH,
        REPORT_PATH,
        *[VISUALS_DIR / f"{number:02d}_{name}.png" for number, name in [
            (1, "annual_performance"),
            (2, "monthly_sales_trend"),
            (3, "category_performance"),
            (4, "subcategory_profit"),
            (5, "region_performance"),
            (6, "discount_profitability"),
            (7, "state_profit_extremes"),
            (8, "top_customers"),
        ]],
    ]
    missing = [str(path) for path in expected_files if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing analysis outputs: {', '.join(missing)}")

    total_sales = kpis.loc[kpis["KPI"].eq("Total Sales"), "Value"].iloc[0]
    category_sales = tables["Categories"]["Sales"].sum()
    if not np.isclose(total_sales, category_sales):
        raise ValueError("Category sales do not reconcile with total sales.")

    total_profit = kpis.loc[kpis["KPI"].eq("Total Profit"), "Value"].iloc[0]
    region_profit = tables["Regions"]["Profit"].sum()
    if not np.isclose(total_profit, region_profit):
        raise ValueError("Regional profit does not reconcile with total profit.")


def main() -> None:
    """Run the complete business analysis pipeline."""
    data = load_data()
    tables = build_analysis_tables(data)
    kpis = build_kpis(data, tables)
    save_outputs(data, kpis, tables)
    verify_outputs(kpis, tables)

    print(f"Rows analyzed: {len(data):,}")
    print(f"KPIs created: {len(kpis):,}")
    print(f"Report: {REPORT_PATH}")
    print(f"Excel workbook: {EXCEL_PATH}")
    print(f"Charts: {VISUALS_DIR}")


if __name__ == "__main__":
    main()
